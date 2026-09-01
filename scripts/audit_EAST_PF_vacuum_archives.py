#!/usr/bin/env python3
"""Compare the historical EAST PF archive with regenerated GL32 fields."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys

import matplotlib.pyplot as plt
from matplotlib.colors import SymLogNorm
import numpy as np
from openpyxl import load_workbook

REPOSITORY = Path(__file__).resolve().parents[1]
if str(REPOSITORY) not in sys.path:
    sys.path.insert(0, str(REPOSITORY))

from pyna.toroidal.coils import (  # noqa: E402
    BRBZ_induced_by_rectangular_winding_pack_gauss_legendre,
)


DEFAULT_LEGACY_ROOT = Path(
    "/mnt/c/Users/Wenyin/OneDrive/MCFdata/EAST/allcoils/"
    "PF_legacy_nonconverged_20260901"
)
DEFAULT_FRESH_ROOT = Path(
    "/mnt/c/Users/Wenyin/OneDrive/MCFdata/EAST/allcoils/PF"
)
DEFAULT_OUT_DIR = Path(
    "/home/wenyin/repos/pyna/artifacts/"
    "east_pf_gl32_regeneration_20260901/audit"
)


def _load_coils(workbook: Path) -> dict[str, tuple[float, float, float, float, int]]:
    rows = list(
        load_workbook(workbook, data_only=True, read_only=True)
        .active.iter_rows(values_only=True)
    )
    header = {str(value).strip().lower(): index for index, value in enumerate(rows[0])}
    return {
        str(row[header["label"]]): (
            float(row[header["rc"]]),
            float(row[header["zc"]]),
            float(row[header["width"]]),
            float(row[header["height"]]),
            int(row[header["turn"]]),
        )
        for row in rows[1:]
    }


def _load_field(path: Path) -> dict[str, np.ndarray]:
    with np.load(path, allow_pickle=False) as archive:
        return {
            key: np.asarray(archive[key], dtype=np.float64)
            for key in ("R", "Z", "BR", "BZ")
        }


def _relative_l2(first: np.ndarray, second: np.ndarray) -> float:
    return float(np.linalg.norm(first - second) / np.linalg.norm(second))


def run(
    *,
    legacy_root: Path,
    fresh_root: Path,
    out_dir: Path,
    workbook: Path,
    low_order: int,
    max_workers: int,
    backend: str,
) -> tuple[Path, Path]:
    coils = _load_coils(workbook)
    legacy_fields: dict[str, dict[str, np.ndarray]] = {}
    fresh_fields: dict[str, dict[str, np.ndarray]] = {}
    per_coil: dict[str, dict[str, float | int | str]] = {}

    for number in range(1, 15):
        label = f"PF{number}"
        filename = f"EAST_{label}_1kA.npz"
        legacy = _load_field(legacy_root / filename)
        fresh = _load_field(fresh_root / filename)
        if not np.array_equal(legacy["R"], fresh["R"]) or not np.array_equal(
            legacy["Z"], fresh["Z"]
        ):
            raise ValueError(f"{label}: legacy and fresh grids differ")

        R, Z = fresh["R"], fresh["Z"]
        RR, ZZ = np.meshgrid(R, Z, indexing="ij")
        low_BR, low_BZ = BRBZ_induced_by_rectangular_winding_pack_gauss_legendre(
            *coils[label],
            1000.0,
            RR,
            ZZ,
            quadrature_order=low_order,
            max_workers=max_workers,
            backend=backend,
        )
        legacy_vector = np.stack([legacy["BR"], legacy["BZ"]])
        fresh_vector = np.stack([fresh["BR"], fresh["BZ"]])
        low_vector = np.stack([low_BR, low_BZ])
        delta_BZ = fresh["BZ"] - legacy["BZ"]
        per_coil[label] = {
            "filename": filename,
            "relative_l2_legacy_to_fresh": _relative_l2(
                legacy_vector, fresh_vector
            ),
            f"relative_l2_GL{low_order}_to_fresh": _relative_l2(
                low_vector, fresh_vector
            ),
            "max_abs_delta_BR_T": float(
                np.max(np.abs(fresh["BR"] - legacy["BR"]))
            ),
            "max_abs_delta_BZ_T": float(np.max(np.abs(delta_BZ))),
            "fresh_vector_rms_T": float(
                np.sqrt(np.mean(fresh["BR"] ** 2 + fresh["BZ"] ** 2))
            ),
        }
        legacy_fields[label] = legacy
        fresh_fields[label] = fresh

    mirror_pairs: dict[str, float] = {}
    for upper_number in range(1, 15, 2):
        upper_label = f"PF{upper_number}"
        lower_label = f"PF{upper_number + 1}"
        upper = fresh_fields[upper_label]
        lower = fresh_fields[lower_label]
        mirror_error = np.stack(
            [
                upper["BR"] + lower["BR"][:, ::-1],
                upper["BZ"] - lower["BZ"][:, ::-1],
            ]
        )
        mirror_reference = np.stack([upper["BR"], upper["BZ"]])
        pair_name = f"{upper_label}/{lower_label}"
        mirror_pairs[pair_name] = float(
            np.linalg.norm(mirror_error) / np.linalg.norm(mirror_reference)
        )

    out_dir.mkdir(parents=True, exist_ok=True)
    figure_path = out_dir / "EAST_PF_legacy_vs_GL32_delta_BZ_2x7.png"
    all_delta_BZ_mT = [
        1.0e3
        * (fresh_fields[f"PF{number}"]["BZ"] - legacy_fields[f"PF{number}"]["BZ"])
        for number in range(1, 15)
    ]
    global_limit = max(float(np.max(np.abs(value))) for value in all_delta_BZ_mT)
    norm = SymLogNorm(
        linthresh=global_limit * 1.0e-7,
        linscale=0.8,
        vmin=-global_limit,
        vmax=global_limit,
        base=10,
    )
    fig, axes = plt.subplots(
        2,
        7,
        figsize=(15.0, 4.8),
        sharex=True,
        sharey=True,
        constrained_layout=True,
    )
    image = None
    for column, upper_number in enumerate(range(1, 15, 2)):
        labels = (f"PF{upper_number}", f"PF{upper_number + 1}")
        delta_pair = [
            1.0e3
            * (fresh_fields[label]["BZ"] - legacy_fields[label]["BZ"])
            for label in labels
        ]
        pair_limit = max(float(np.max(np.abs(value))) for value in delta_pair)
        for row, (label, delta_BZ) in enumerate(zip(labels, delta_pair, strict=True)):
            axis = axes[row, column]
            field = fresh_fields[label]
            image = axis.pcolormesh(
                field["R"],
                field["Z"],
                delta_BZ.T,
                shading="auto",
                cmap="coolwarm",
                norm=norm,
            )
            metrics = per_coil[label]
            axis.set_title(
                f"{label}  relL2={metrics['relative_l2_legacy_to_fresh']:.2e}",
                fontsize=7.8,
                pad=2,
            )
            axis.set_aspect("equal")
            axis.tick_params(labelsize=6.5, length=2, pad=1)
            if column == 0:
                axis.tick_params(labelleft=True)
        pair_name = f"{labels[0]}/{labels[1]}"
        axes[1, column].text(
            0.5,
            0.03,
            f"pair max={pair_limit:.2e} mT\nmirror={mirror_pairs[pair_name]:.1e}",
            transform=axes[1, column].transAxes,
            ha="center",
            va="bottom",
            fontsize=5.8,
            color="black",
            bbox={"facecolor": "white", "alpha": 0.75, "edgecolor": "none", "pad": 1.2},
        )

    fig.suptitle(
        "EAST PF vacuum fields: regenerated GL32 minus legacy BZ",
        fontsize=10,
    )
    fig.supxlabel("R [m]", fontsize=8)
    fig.supylabel("Z [m]", fontsize=8)
    colorbar = fig.colorbar(image, ax=axes.ravel().tolist(), fraction=0.014, pad=0.01)
    colorbar.set_label("ΔBZ fresh−legacy [mT]", fontsize=7.5)
    colorbar.ax.tick_params(labelsize=6, length=2)
    fig.savefig(figure_path, dpi=200)
    plt.close(fig)

    report_path = out_dir / "EAST_PF_legacy_vs_GL32_audit.json"
    report = {
        "schema": "pyna.EAST_PF_vacuum_archive_comparison.v1",
        "legacy_root": str(legacy_root),
        "fresh_root": str(fresh_root),
        "workbook": str(workbook),
        "fresh_quadrature_order": 32,
        "convergence_comparison_order": int(low_order),
        "array_indexing": "(R,Z)",
        "per_coil": per_coil,
        "fresh_up_down_mirror_relative_l2": mirror_pairs,
    }
    report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    return figure_path, report_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--legacy-root", type=Path, default=DEFAULT_LEGACY_ROOT)
    parser.add_argument("--fresh-root", type=Path, default=DEFAULT_FRESH_ROOT)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--low-order", type=int, default=16)
    parser.add_argument("--max-workers", type=int, default=8)
    parser.add_argument("--backend", choices=("auto", "cpu", "cuda"), default="auto")
    args = parser.parse_args()
    workbook = args.workbook or args.legacy_root / "EAST_PF_coils.xlsx"
    figure_path, report_path = run(
        legacy_root=args.legacy_root,
        fresh_root=args.fresh_root,
        out_dir=args.out_dir,
        workbook=workbook,
        low_order=args.low_order,
        max_workers=args.max_workers,
        backend=args.backend,
    )
    print(figure_path)
    print(report_path)


if __name__ == "__main__":
    main()
