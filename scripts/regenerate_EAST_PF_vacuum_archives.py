#!/usr/bin/env python3
"""Regenerate the 14 EAST PF vacuum-field archives with the GL32 loop kernel.

The default mode writes a complete four-key ``R/Z/BR/BZ`` candidate set to a
staging directory.  The authoritative OneDrive files are touched only when
``--replace-authority`` is supplied explicitly.
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import time
from typing import Any

import numpy as np


REPOSITORY = Path(__file__).resolve().parents[1]
if str(REPOSITORY) not in sys.path:
    sys.path.insert(0, str(REPOSITORY))

from pyna.toroidal.coils import (  # noqa: E402
    BRBZ_induced_by_rectangular_winding_pack_gauss_legendre,
)


DEFAULT_ARCHIVE_DIR = Path(
    "/mnt/c/Users/Wenyin/OneDrive/MCFdata/EAST/allcoils/PF"
)
BACKUP_DIRECTORY_NAME = "PF_legacy_nonconverged_20260901"
PF_LABELS = tuple(f"PF{index}" for index in range(1, 15))
CURRENT_PER_TURN_A = 1000.0
HIGH_ORDER = 32
LOW_ORDER = 16
REPORT_NAME = "EAST_PF_vacuum_GL32_regeneration_audit.json"


def _load_geometry(workbook: Path) -> dict[str, dict[str, float | int]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to read EAST_PF_coils.xlsx; install it in "
            "the Python environment used for this regeneration script"
        ) from exc

    book = load_workbook(workbook, data_only=True, read_only=True)
    rows = list(book.active.iter_rows(values_only=True))
    header = {
        str(value).strip().lower(): index for index, value in enumerate(rows[0])
    }
    required = ("label", "rc", "zc", "width", "height", "turn")
    missing = [name for name in required if name not in header]
    if missing:
        raise ValueError(f"PF workbook is missing columns {missing}: {workbook}")

    geometry: dict[str, dict[str, float | int]] = {}
    for row in rows[1:]:
        label = str(row[header["label"]]).strip().upper()
        if label not in PF_LABELS:
            continue
        geometry[label] = {
            "center_R_m": float(row[header["rc"]]),
            "center_Z_m": float(row[header["zc"]]),
            "width_m": float(row[header["width"]]),
            "height_m": float(row[header["height"]]),
            "turns": int(row[header["turn"]]),
        }
    missing_labels = [label for label in PF_LABELS if label not in geometry]
    if missing_labels:
        raise ValueError(f"PF workbook is missing coils {missing_labels}: {workbook}")
    return geometry


def _load_legacy_archive(
    path: Path,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    with np.load(path, allow_pickle=False) as archive:
        if set(archive.files) != {"R", "Z", "BR", "BZ"}:
            raise ValueError(
                f"expected the four-key R/Z/BR/BZ archive format: {path}"
            )
        R = np.asarray(archive["R"], dtype=np.float64)
        Z = np.asarray(archive["Z"], dtype=np.float64)
        BR = np.asarray(archive["BR"], dtype=np.float64)
        BZ = np.asarray(archive["BZ"], dtype=np.float64)
    if R.ndim != 1 or Z.ndim != 1 or BR.shape != (R.size, Z.size):
        raise ValueError(f"invalid EAST PF archive grid or BR shape: {path}")
    if BZ.shape != BR.shape:
        raise ValueError(f"BR/BZ shapes differ: {path}")
    return R, Z, BR, BZ


def _write_four_key_npz(
    path: Path,
    *,
    R: np.ndarray,
    Z: np.ndarray,
    BR: np.ndarray,
    BZ: np.ndarray,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.stem}.", suffix=".npz", dir=path.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        np.savez(temporary, R=R, Z=Z, BR=BR, BZ=BZ)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _relative_l2(left: np.ndarray, right: np.ndarray) -> float:
    return float(np.linalg.norm(left - right) / np.linalg.norm(right))


def _resolved_backend(requested: str) -> str:
    if requested == "cpu":
        return "cpu"
    try:
        from pyna.toroidal.coils import accel

        available = bool(accel._CUPY_AVAILABLE)
        if available:
            available = accel.cp.cuda.runtime.getDeviceCount() > 0
    except (ImportError, RuntimeError):
        available = False
    if requested == "cuda" and not available:
        raise RuntimeError("--backend cuda requires CuPy and a visible CUDA device")
    return "cuda" if available else "cpu"


def _regenerate(
    *,
    archive_dir: Path,
    workbook: Path,
    output_dir: Path,
    backend: str,
    max_workers: int | None,
) -> dict[str, Any]:
    geometry = _load_geometry(workbook)
    resolved_backend = _resolved_backend(backend)
    coil_reports: list[dict[str, Any]] = []
    started = time.perf_counter()

    for label in PF_LABELS:
        filename = f"EAST_{label}_1kA.npz"
        legacy_path = archive_dir / filename
        R, Z, legacy_BR, legacy_BZ = _load_legacy_archive(legacy_path)
        RR, ZZ = np.meshgrid(R, Z, indexing="ij")
        coil = geometry[label]
        arguments = (
            coil["center_R_m"],
            coil["center_Z_m"],
            coil["width_m"],
            coil["height_m"],
            coil["turns"],
            CURRENT_PER_TURN_A,
            RR,
            ZZ,
        )

        high_started = time.perf_counter()
        fresh_BR, fresh_BZ = (
            BRBZ_induced_by_rectangular_winding_pack_gauss_legendre(
                *arguments,
                quadrature_order=HIGH_ORDER,
                max_workers=max_workers,
                backend=backend,
            )
        )
        high_seconds = time.perf_counter() - high_started

        low_started = time.perf_counter()
        low_BR, low_BZ = (
            BRBZ_induced_by_rectangular_winding_pack_gauss_legendre(
                *arguments,
                quadrature_order=LOW_ORDER,
                max_workers=max_workers,
                backend=backend,
            )
        )
        low_seconds = time.perf_counter() - low_started

        fresh = np.stack([fresh_BR, fresh_BZ])
        lower = np.stack([low_BR, low_BZ])
        legacy = np.stack([legacy_BR, legacy_BZ])
        staged_path = output_dir / filename
        _write_four_key_npz(
            staged_path,
            R=R,
            Z=Z,
            BR=fresh_BR,
            BZ=fresh_BZ,
        )
        coil_report = {
            "label": label,
            "filename": filename,
            "geometry": {
                **coil,
                "current_per_turn_A": CURRENT_PER_TURN_A,
                "total_ampere_turns_A": (
                    int(coil["turns"]) * CURRENT_PER_TURN_A
                ),
            },
            "grid_shape": [int(R.size), int(Z.size)],
            "backend_requested": backend,
            "backend_resolved": resolved_backend,
            "GL32_seconds": high_seconds,
            "GL16_seconds": low_seconds,
            "GL16_to_GL32_relative_L2": _relative_l2(lower, fresh),
            "legacy_to_GL32_relative_L2": _relative_l2(fresh, legacy),
            "legacy_GL32_cosine_similarity": float(
                np.vdot(fresh, legacy).real
                / (np.linalg.norm(fresh) * np.linalg.norm(legacy))
            ),
        }
        coil_reports.append(coil_report)
        print(
            f"{label}: {resolved_backend} GL32 {high_seconds:.3f} s; "
            f"GL16->32 {coil_report['GL16_to_GL32_relative_L2']:.3e}; "
            f"legacy->GL32 {coil_report['legacy_to_GL32_relative_L2']:.3e}",
            flush=True,
        )

    return {
        "schema": "pyna.EAST_PF_vacuum_GL32_regeneration_audit.v1",
        "archive_directory": str(archive_dir),
        "workbook": str(workbook),
        "staging_directory": str(output_dir),
        "archive_format_keys": ["R", "Z", "BR", "BZ"],
        "current_per_turn_A": CURRENT_PER_TURN_A,
        "quadrature_orders": [LOW_ORDER, HIGH_ORDER],
        "backend_requested": backend,
        "backend_resolved": resolved_backend,
        "total_generation_seconds": time.perf_counter() - started,
        "authority_replaced": False,
        "legacy_backup_directory": None,
        "coils": coil_reports,
    }


def _replace_authority(
    *, archive_dir: Path, output_dir: Path
) -> Path:
    backup_dir = archive_dir.parent / BACKUP_DIRECTORY_NAME
    if backup_dir.exists():
        raise FileExistsError(
            f"refusing to overwrite the existing legacy backup: {backup_dir}"
        )

    prepared: dict[str, Path] = {}
    for label in PF_LABELS:
        filename = f"EAST_{label}_1kA.npz"
        staged = output_dir / filename
        _load_legacy_archive(staged)
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{filename}.GL32.", suffix=".npz", dir=archive_dir
        )
        os.close(descriptor)
        temporary = Path(temporary_name)
        shutil.copyfile(staged, temporary)
        prepared[filename] = temporary

    try:
        backup_dir.mkdir()
        for label in PF_LABELS:
            filename = f"EAST_{label}_1kA.npz"
            shutil.copy2(archive_dir / filename, backup_dir / filename)
        shutil.copy2(
            archive_dir / "EAST_PF_coils.xlsx",
            backup_dir / "EAST_PF_coils.xlsx",
        )
        for label in PF_LABELS:
            filename = f"EAST_{label}_1kA.npz"
            os.replace(prepared[filename], archive_dir / filename)
    finally:
        for temporary in prepared.values():
            temporary.unlink(missing_ok=True)
    return backup_dir


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--archive-dir",
        type=Path,
        default=DEFAULT_ARCHIVE_DIR,
        help="directory containing EAST_PF1_1kA.npz through EAST_PF14_1kA.npz",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=None,
        help="PF geometry workbook; default is EAST_PF_coils.xlsx in archive-dir",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        required=True,
        help="staging directory for the regenerated four-key NPZ files",
    )
    parser.add_argument(
        "--backend", choices=("auto", "cpu", "cuda"), default="auto"
    )
    parser.add_argument("--max-workers", type=int, default=None)
    parser.add_argument(
        "--replace-authority",
        action="store_true",
        help=(
            "after all 14 staging files succeed, copy the old files to "
            f"../{BACKUP_DIRECTORY_NAME} and atomically replace the authority files"
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    archive_dir = args.archive_dir.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    workbook = (
        args.workbook.expanduser().resolve()
        if args.workbook is not None
        else archive_dir / "EAST_PF_coils.xlsx"
    )
    if output_dir == archive_dir:
        raise ValueError("--output-dir must be a separate staging directory")
    output_dir.mkdir(parents=True, exist_ok=True)

    report = _regenerate(
        archive_dir=archive_dir,
        workbook=workbook,
        output_dir=output_dir,
        backend=args.backend,
        max_workers=args.max_workers,
    )
    if args.replace_authority:
        backup_dir = _replace_authority(
            archive_dir=archive_dir,
            output_dir=output_dir,
        )
        report["authority_replaced"] = True
        report["legacy_backup_directory"] = str(backup_dir)

    report_path = output_dir / REPORT_NAME
    report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(f"audit: {report_path}")
    if not args.replace_authority:
        print("authority unchanged (staging-only mode)")


if __name__ == "__main__":
    main()
